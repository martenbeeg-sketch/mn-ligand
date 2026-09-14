from __future__ import annotations

import io
import json
from pathlib import Path
from unittest.mock import patch

import pandas as pd
import pytest
from openpyxl import Workbook
from streamlit.testing.v1 import AppTest

from mn_ligand.workflows.compound_preparation import (
    annotation_stripped_smiles_candidate,
    annotate_component_relationships,
    analyze_tabular_compound_dataset,
    compound_dataset_columns,
    compound_dataset_sheets,
    compound_box_fit_rows,
    create_compound_import_job,
    create_docking_parent_selection_job,
    docking_parent_record,
    parent_duplicate_report,
    summarize_compound_dataset,
)
from mn_ligand.workflows.docking import load_compound_records
from mn_ligand.workflows.compound_pubchem import (
    create_compound_review_job,
    pubchem_component_options,
    selected_parent_candidate,
)


SDF_DATA = b"""mol-1
  test

  0  0  0  0  0  0  0  0  0  0999 V2000
M  END
$$$$
mol-2
  test

  0  0  0  0  0  0  0  0  0  0999 V2000
M  END
$$$$
"""


def test_compound_dataset_formats_are_counted() -> None:
    assert summarize_compound_dataset(SDF_DATA, "library.sdf")["compound_count"] == 2
    assert summarize_compound_dataset(b"CCO ethanol\nCCN ethylamine\n", "library.smi")["compound_count"] == 2
    assert summarize_compound_dataset(
        b"name,smiles\nethanol,CCO\nethylamine,CCN\n", "library.csv"
    )["compound_count"] == 2


def test_csv_requires_smiles_column() -> None:
    with pytest.raises(ValueError, match="SMILES column"):
        summarize_compound_dataset(b"name,value\nethanol,1\n", "library.csv")


def test_csv_column_mapping_publishes_canonical_compound_set(tmp_path: Path) -> None:
    source = b"catalog_key,structure,activity\nA-1,CCO,7.1\nA-2,CCN,6.5\n"
    with patch.dict("os.environ", {"MN_LIGAND_RUN_DIR": str(tmp_path)}, clear=False):
        job = create_compound_import_job(
            source,
            filename="mapped.csv",
            id_column="catalog_key",
            smiles_column="structure",
        )

    assert job.artifact_manifest is not None
    normalized = job.artifact_manifest.by_type("compound_set")[0]
    source_artifact = job.artifact_manifest.by_type("source_compound_dataset")[0]
    normalized_path = normalized.resolve(job.run_dir, must_exist=True)
    assert normalized_path is not None
    frame = pd.read_csv(normalized_path)
    assert frame[["compound_id", "smiles", "activity"]].to_dict("records") == [
        {"compound_id": "A-1", "smiles": "CCO", "activity": 7.1},
        {"compound_id": "A-2", "smiles": "CCN", "activity": 6.5},
    ]
    assert {
        "formula",
        "molecular_weight",
        "clogp",
        "tpsa",
        "qed",
    }.issubset(frame.columns)
    assert source_artifact.resolve(job.run_dir, must_exist=True).read_bytes() == source


def test_excel_import_preserves_workbook_and_rejects_invalid_smiles(
    tmp_path: Path,
) -> None:
    workbook = Workbook()
    compounds = workbook.active
    compounds.title = "Compounds"
    compounds.append(["Catalog ID", "Structure", "Assay", "Comment"])
    compounds.append(["A-1", "CCO", 7.1, "active"])
    compounds.append(["A-2", "not-a-smiles", 6.5, "review"])
    compounds.append(["A-3", "c1ccccc1", 5.0, "reference"])
    notes = workbook.create_sheet("Metadata")
    notes.append(["Project", "Example screen"])
    source = io.BytesIO()
    workbook.save(source)
    source_bytes = source.getvalue()

    assert compound_dataset_sheets(source_bytes, "library.xlsx") == (
        "Compounds",
        "Metadata",
    )
    assert compound_dataset_columns(
        source_bytes,
        "library.xlsx",
        sheet_name="Compounds",
    ) == ("Catalog ID", "Structure", "Assay", "Comment")
    analysis = analyze_tabular_compound_dataset(
        source_bytes,
        "library.xlsx",
        sheet_name="Compounds",
        id_column="Catalog ID",
        smiles_column="Structure",
    )
    assert analysis["summary"]["valid_count"] == 2
    assert analysis["summary"]["invalid_count"] == 1
    assert analysis["invalid_rows"][0]["Comment"] == "review"

    with patch.dict(
        "os.environ", {"MN_LIGAND_RUN_DIR": str(tmp_path)}, clear=False
    ):
        job = create_compound_import_job(
            source_bytes,
            filename="library.xlsx",
            dataset_name="Excel screen",
            sheet_name="Compounds",
            id_column="Catalog ID",
            smiles_column="Structure",
        )

    assert job.metadata["compound_count"] == 2
    assert job.metadata["invalid_compound_count"] == 1
    assert job.artifact_manifest is not None
    original = job.artifact_manifest.by_type("source_compound_dataset")[0]
    assert original.resolve(job.run_dir, must_exist=True).read_bytes() == source_bytes
    assert job.artifact_manifest.by_type("compound_validation_report")
    rejected = job.artifact_manifest.by_type("rejected_compounds")[0]
    rejected_frame = pd.read_csv(rejected.resolve(job.run_dir, must_exist=True))
    assert rejected_frame.loc[0, "compound_id"] == "A-2"
    assert rejected_frame.loc[0, "Comment"] == "review"
    report = json.loads(
        job.artifact_manifest.by_type("import_report")[0]
        .resolve(job.run_dir, must_exist=True)
        .read_text()
    )
    assert report["workbook_sheets"] == ["Compounds", "Metadata"]
    normalized = job.artifact_manifest.by_type("compound_set")[0]
    normalized_path = normalized.resolve(job.run_dir, must_exist=True)
    assert load_compound_records([normalized_path]) == [
        {"compound_id": "A-1", "smiles": "CCO"},
        {"compound_id": "A-3", "smiles": "c1ccccc1"},
    ]


def test_excel_reports_components_and_vendor_annotation_candidates() -> None:
    workbook = Workbook()
    compounds = workbook.active
    compounds.title = "Compound Information"
    compounds.append(["Catalog Number", "SMILES", "CAS Number"])
    compounds.append(["A-1", "CC[NH3+].[Cl-]", ""])
    compounds.append(["A-2", "CCO.[Z]", "64-17-5"])
    source = io.BytesIO()
    workbook.save(source)

    analysis = analyze_tabular_compound_dataset(
        source.getvalue(),
        "library.xlsx",
        sheet_name="Compound Information",
        id_column="Catalog Number",
        smiles_column="SMILES",
    )

    assert analysis["summary"]["valid_count"] == 1
    assert analysis["summary"]["invalid_count"] == 1
    assert analysis["summary"]["multi_fragment_count"] == 1
    assert analysis["summary"]["charged_component_count"] == 2
    assert len(analysis["component_rows"]) == 2
    assert analysis["invalid_rows"][0]["annotation_stripped_candidate"] == "CCO"
    assert analysis["invalid_rows"][0]["candidate_status"].startswith(
        "Review required"
    )


@pytest.mark.parametrize(
    ("source_smiles", "expected"),
    [
        ("CCO.O.[1/4]", "CCO.O"),
        ("CCN.[1.5 CH3COOH]", "CCN"),
    ],
)
def test_vendor_fraction_annotations_are_removed_for_review_preview(
    source_smiles: str,
    expected: str,
) -> None:
    assert annotation_stripped_smiles_candidate(source_smiles) == expected


def test_parent_duplicate_report_includes_salts_and_repeated_parents() -> None:
    report = parent_duplicate_report(
        [
            {"compound_id": "free", "smiles": "CN"},
            {"compound_id": "salt", "smiles": "C[NH3+].[Cl-]"},
            {"compound_id": "repeat", "smiles": "CN.CN.[Cl-]"},
            {"compound_id": "other", "smiles": "CCO"},
            {"compound_id": "ambiguous", "smiles": "CC.CN"},
            {"compound_id": "stereo-r", "smiles": "C[C@H](O)F"},
            {"compound_id": "stereo-s", "smiles": "C[C@@H](O)F"},
        ]
    )

    assert report["summary"] == {
        "unique_parent_count": 4,
        "duplicate_group_count": 1,
        "duplicate_entry_count": 3,
        "redundant_entry_count": 2,
        "ambiguous_parent_count": 1,
    }
    duplicate_ids = {row["compound_id"] for row in report["rows"]}
    assert duplicate_ids == {"free", "salt", "repeat"}
    repeated = next(
        row for row in report["rows"] if row["compound_id"] == "repeat"
    )
    assert repeated["parent_occurrences"] == 2
    assert repeated["source_fragment_count"] == 3
    assert report["rows"][0]["standardized_parent_smiles"] == "CN"
    parent_rows = report["docking_parent_rows"]
    assert len(parent_rows) == 4
    methylamine = next(
        row
        for row in parent_rows
        if row["standardized_parent_smiles"] == "CN"
    )
    assert methylamine["source_record_count"] == 3
    assert set(methylamine["compound_ids"].split(" | ")) == {
        "free",
        "salt",
        "repeat",
    }
    ambiguous = docking_parent_record("CC.CN")
    assert ambiguous["parent_status"].startswith("ambiguous")


def test_docking_parent_normalizes_bonded_metals_and_isotopic_hydrogens() -> None:
    sodium_salt = docking_parent_record(
        "Cc1cc(OCP(=O)(O)[O][Na])cc(C)c1Cc1ccc(O)c(C(C)C)c1"
    )
    assert sodium_salt["source_fragment_count"] == 1
    assert sodium_salt["docking_fragment_count"] == 2
    assert "Na" not in sodium_salt["standardized_parent_smiles"]
    assert sodium_salt["standardized_parent_smiles"] == (
        "Cc1cc(OCP(=O)(O)O)cc(C)c1Cc1ccc(O)c(C(C)C)c1"
    )

    deuterated = docking_parent_record(
        "[2H]C([2H])([2H])C(c1cc(Oc2c(Cl)cc(-n3nc(C#N)c(=O)[nH]c3=O)"
        "cc2Br)n[nH]c1=O)C([2H])([2H])[2H]"
    )
    assert "[2H]" not in deuterated["standardized_parent_smiles"]
    assert deuterated["standardized_parent_smiles"].startswith("CC(C)c1")


def test_docking_parent_selection_publishes_typed_subset(
    tmp_path: Path,
) -> None:
    source = b"id,smiles,name\nfree,CN,Parent\nsalt,C[NH3+].[Cl-],Salt\nother,CCO,Other\n"
    with patch.dict(
        "os.environ", {"MN_LIGAND_RUN_DIR": str(tmp_path)}, clear=False
    ):
        imported = create_compound_import_job(
            source,
            filename="parents.csv",
            id_column="id",
            smiles_column="smiles",
        )
        source_artifact = imported.artifact_manifest.by_type(
            "compound_set"
        )[0]
        source_path = source_artifact.resolve(
            imported.run_dir, must_exist=True
        )
        source_rows = pd.read_csv(source_path).fillna("").to_dict("records")
        parents = parent_duplicate_report(source_rows)[
            "docking_parent_rows"
        ]
        parents[0]["modeling_smiles"] = "C[NH3+]"
        parents[0]["modeling_preparation"] = "Scrub pH 7.4"
        parents[0]["modeling_ph"] = 7.4
        selected = create_docking_parent_selection_job(
            source_job=imported,
            source_artifact=source_artifact,
            parent_rows=parents[:1],
            selection_mode="Manual selection",
            excluded_parent_rows=[
                {
                    **parents[1],
                    "estimated_3d_length_angstrom": 34.2,
                    "estimated_3d_width_angstrom": 12.0,
                    "estimated_3d_thickness_angstrom": 5.7,
                    "estimated_max_span_angstrom": 34.4,
                    "box_fit_reason": "length 34.200 Å > box 30.000 Å",
                    "box_fit_max_excess_angstrom": 4.2,
                }
            ],
            box_size=(30.0, 16.0, 16.0),
        )

    selected_artifact = selected.artifact_manifest.by_type("compound_set")[0]
    selected_path = selected_artifact.resolve(
        selected.run_dir, must_exist=True
    )
    selected_frame = pd.read_csv(selected_path)
    assert len(selected_frame) == 1
    assert {
        "compound_id",
        "smiles",
        "docking_parent_id",
        "source_compound_ids",
        "identity_parent_smiles",
        "modeling_smiles",
        "modeling_formal_charge",
    }.issubset(selected_frame.columns)
    assert selected_frame.iloc[0]["identity_parent_smiles"] == "CN"
    assert selected_frame.iloc[0]["smiles"] == "C[NH3+]"
    assert selected_frame.iloc[0]["modeling_smiles"] == "C[NH3+]"
    assert selected_frame.iloc[0]["modeling_formal_charge"] == 1
    assert selected.metadata["parent_run_id"] == imported.run_id
    assert selected.metadata["selection_mode"] == "Manual selection"
    assert selected.metadata["modeling_state_prepared"] is True
    assert selected_artifact.metadata["modeling_state_prepared"] is True
    assert selected.metadata["excluded_compound_count"] == 1
    assert selected.metadata["docking_box_size_angstrom"] == [
        30.0,
        16.0,
        16.0,
    ]
    exclusion_artifact = selected.artifact_manifest.by_type(
        "compound_exclusion_report"
    )[0]
    exclusion_frame = pd.read_csv(
        exclusion_artifact.resolve(selected.run_dir, must_exist=True)
    )
    assert len(exclusion_frame) == 1
    assert exclusion_frame.iloc[0]["exclusion_category"] == "docking_box_fit"
    assert (
        exclusion_frame.iloc[0]["exclusion_reason"]
        == "length 34.200 Å > box 30.000 Å"
    )
    assert exclusion_frame.iloc[0]["box_size_x_angstrom"] == 30.0
    selection_input = json.loads((selected.run_dir / "input.json").read_text())
    assert selection_input["excluded_docking_parent_ids"] == [
        str(parents[1]["docking_parent_id"])
    ]
    assert load_compound_records([selected_path]) == [
        {
            "compound_id": str(selected_frame.iloc[0]["compound_id"]),
            "smiles": str(selected_frame.iloc[0]["smiles"]),
        }
    ]


def test_compound_box_fit_preflight_uses_full_box_principal_dimensions(
    tmp_path: Path, monkeypatch
) -> None:
    monkeypatch.setenv("MN_LIGAND_RUN_DIR", str(tmp_path / "runs"))
    rows = compound_box_fit_rows(
        [
            {
                "docking_parent_id": "small",
                "standardized_parent_smiles": "CCO",
            },
            {
                "docking_parent_id": "large",
                "standardized_parent_smiles": "CCCCCCCCCCCCCCCCCCCC",
            },
        ],
        box_size=(10.0, 10.0, 10.0),
    )

    assert rows[0]["box_fit_status"] == "fits estimated box"
    assert rows[1]["box_fit_status"] == "likely too large"
    assert rows[1]["estimated_3d_length_angstrom"] > 18.0
    assert rows[1]["estimated_max_span_angstrom"] >= rows[1][
        "estimated_3d_length_angstrom"
    ]
    assert "length" in rows[1]["box_fit_reason"]
    assert "> box 10.000 Å" in rows[1]["box_fit_reason"]
    assert rows[1]["box_fit_max_excess_angstrom"] > 8.0
    assert "box_fit_clearance_angstrom" not in rows[1]
    cache_path = (
        tmp_path / "cache" / "compound-size-estimates-v1.json"
    )
    cache = json.loads(cache_path.read_text())
    assert cache["entry_count"] == 2
    assert cache["method"] == "etkdg-v3-mmff-uff-principal-axes-v1"

    with patch(
        "mn_ligand.workflows.compound_preparation.estimate_parent_3d_size",
        side_effect=AssertionError("persistent estimates should be reused"),
    ):
        reused = compound_box_fit_rows(
            [
                {
                    "docking_parent_id": "small",
                    "standardized_parent_smiles": "CCO",
                },
                {
                    "docking_parent_id": "large",
                    "standardized_parent_smiles": "CCCCCCCCCCCCCCCCCCCC",
                },
            ],
            box_size=(30.0, 16.0, 16.0),
        )
    assert all(row["box_fit_status"] == "fits estimated box" for row in reused)


def test_components_recognize_counterions_and_find_unformulated_parent() -> None:
    valid_rows = [
        {
            "compound_id": "parent",
            "smiles": "CCN",
            "fragment_count": 1,
        },
        {
            "compound_id": "hydrochloride",
            "smiles": "CCN.Cl",
            "fragment_count": 2,
        },
    ]
    components = [
        {
            "compound_id": "hydrochloride",
            "smiles": "CCN",
            "heavy_atoms": 3,
        },
        {
            "compound_id": "hydrochloride",
            "smiles": "Cl",
            "heavy_atoms": 1,
        },
    ]

    annotated = annotate_component_relationships(valid_rows, components)

    parent = next(row for row in annotated if row["smiles"] == "CCN")
    hydrochloride = next(row for row in annotated if row["smiles"] == "Cl")
    assert parent["parent_candidate"] is True
    assert parent["unformulated_library_matches"] == "parent"
    assert parent["parent_match_status"].startswith("exact single-component")
    assert hydrochloride["recognized_component"] == "hydrochloride / HCl"
    assert hydrochloride["component_category"] == "hydrochloride / HCl"
    assert hydrochloride["formulation_category"] == (
        "recognized: hydrochloride / HCl"
    )


def test_components_categorize_organic_mixtures_without_known_formulation() -> None:
    valid_rows = [
        {
            "compound_id": "mixture",
            "smiles": "CCO.CCN",
            "fragment_count": 2,
        }
    ]
    components = [
        {
            "compound_id": "mixture",
            "smiles": "CCO",
            "heavy_atoms": 3,
            "contains_carbon": True,
        },
        {
            "compound_id": "mixture",
            "smiles": "CCN",
            "heavy_atoms": 3,
            "contains_carbon": True,
        },
    ]

    annotated = annotate_component_relationships(valid_rows, components)

    assert {row["component_category"] for row in annotated} == {
        "unrecognized organic co-component; review"
    }
    assert {row["formulation_category"] for row in annotated} == {
        "multiple organic components; mixture/co-drug/homolog review"
    }
    assert {row["parent_match_status"] for row in annotated} == {
        "ambiguous largest components"
    }


def test_compound_import_publishes_portable_compound_set(tmp_path: Path) -> None:
    with patch.dict("os.environ", {"MN_LIGAND_RUN_DIR": str(tmp_path)}, clear=False):
        job = create_compound_import_job(
            SDF_DATA,
            filename="screening library.sdf",
            dataset_name="Primary screen",
        )

    assert job.status == "completed"
    assert job.metadata["compound_count"] == 2
    assert job.artifact_manifest is not None
    compound_sets = job.artifact_manifest.by_type("compound_set")
    assert len(compound_sets) == 1
    assert compound_sets[0].metadata["compound_count"] == 2
    assert not Path(compound_sets[0].path).is_absolute()
    assert compound_sets[0].resolve(job.run_dir, must_exist=True) is not None
    input_payload = json.loads((job.run_dir / "input.json").read_text())
    assert input_payload["dataset_name"] == "Primary screen"


def test_compound_import_ui_has_one_prepare_page() -> None:
    page_root = Path(__file__).parents[1] / "mn_ligand" / "app" / "pages"
    dataset_page = (page_root / "compound_datasets.py").read_text()
    structure_page = (page_root / "structure_preparation.py").read_text()
    assert "compound_dataset_file" in dataset_page
    assert "compound_dataset_file" not in structure_page


def test_compound_dataset_page_inspects_validated_structures(
    tmp_path: Path,
) -> None:
    source = (
        b"catalog_key,structure,activity\n"
        b"A-1,CCO,7.1\n"
        b"A-2,c1ccccc1,6.5\n"
        b"A-3,not-a-smiles,4.0\n"
    )
    with patch.dict(
        "os.environ", {"MN_LIGAND_RUN_DIR": str(tmp_path)}, clear=False
    ):
        create_compound_import_job(
            source,
            filename="screen.csv",
            dataset_name="Primary screen",
            id_column="catalog_key",
            smiles_column="structure",
        )
        page = AppTest.from_file(
            Path(__file__).parents[1]
            / "mn_ligand"
            / "app"
            / "pages"
            / "compound_datasets.py"
        ).run(timeout=20)

    assert not page.exception
    assert [tab.label for tab in page.tabs] == [
        "Input",
        "Columns",
        "Validation",
        "Run",
        "Results",
    ]
    assert next(
        item for item in page.button
        if item.label == "Register validated compound dataset"
    ).disabled is True
    assert any(
        {"Job", "Compare", "Dataset", "Usable", "Rejected"}.issubset(
            frame.value.columns
        )
        for frame in page.dataframe
    )


def test_compound_import_job_page_renders_dataset_report(
    tmp_path: Path,
) -> None:
    source = (
        b"catalog_key,structure,product_name\n"
        b"A-1,CCO,ethanol\n"
        b"A-2,CC[NH3+].[Cl-],ethylammonium chloride\n"
        b"A-3,CCO.[Z],annotated ethanol\n"
    )
    with patch.dict(
        "os.environ", {"MN_LIGAND_RUN_DIR": str(tmp_path)}, clear=False
    ):
        job = create_compound_import_job(
            source,
            filename="screen.csv",
            dataset_name="Compound report test",
            id_column="catalog_key",
            smiles_column="structure",
        )
        full_candidate = {
            "cid": 1,
            "smiles": "CCO.O",
            "molecular_formula": "C2H8O2",
            "query": "annotated ethanol",
            "query_type": "name",
            "pubchem_url": "https://pubchem.ncbi.nlm.nih.gov/compound/1",
        }
        create_compound_review_job(
            job,
            source_row=4,
            compound_id="A-3",
            decision="accepted",
            query="annotated ethanol",
            query_type="name",
            candidate=selected_parent_candidate(
                full_candidate,
                pubchem_component_options(full_candidate["smiles"])[0],
            ),
        )
        page = AppTest.from_file(
            Path(__file__).parents[1]
            / "mn_ligand"
            / "app"
            / "pages"
            / "job_results.py"
        )
        page.query_params["task_group"] = "compound-import"
        page.query_params["run_id"] = job.run_id
        page.run(timeout=20)

    assert not page.exception
    assert "Dataset" in [tab.label for tab in page.tabs]
    assert {
        "Usable compounds",
        "Multi-fragment compounds",
        "Docking-ready parents",
        "Parent duplicates",
        "Rejected / invalid",
        "Profiles",
    }.issubset({tab.label for tab in page.tabs})
    metrics = {metric.label: metric.value for metric in page.metric}
    assert metrics["Usable"] == "3"
    assert metrics["Needs review"] == "0"
    assert metrics["Unique parent SMILES"] == "2"
    assert metrics["Redundant sources"] == "1"
    assert metrics["Multi-component"] == "1"
    assert any(
        getattr(link, "label", "")
        == "Compare target campaigns for this dataset"
        for link in page.get("link_button")
    )
    assert any(
        "PubChem confirmed import" in frame.value.astype(str).to_string()
        for frame in page.dataframe
    )
    assert any(
        {
            "duplicate_group",
            "docking_parent_smiles",
            "standardized_parent_smiles",
        }.issubset(frame.value.columns)
        for frame in page.dataframe
    )
    assert any(
        "annotation_stripped_candidate" in frame.value.columns
        for frame in page.dataframe
    )
    assert any(
        "There are no compounds in this rejected-row view" in item.value
        for item in page.success
    )
    page_source = (
        Path(__file__).parents[1]
        / "mn_ligand"
        / "app"
        / "pages"
        / "compound_results.py"
    ).read_text()
    assert "Skip / keep rejected" in page_source
    assert "Vendor-side structure" in page_source
    assert "Full PubChem record" in page_source
    assert "Parent to import" in page_source
    assert any(
        {"compound_id", "component", "component_count"}.issubset(
            frame.value.columns
        )
        for frame in page.dataframe
    )
    assert any(
        {
            "compound_id",
            "parent_candidate_smiles",
            "other_fragment_smiles",
            "all_fragment_smiles",
        }.issubset(frame.value.columns)
        for frame in page.dataframe
    )
